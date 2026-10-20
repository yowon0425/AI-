"""3단계 B (10/16): 시군 진단표 + 철도 공백 지수 + 시군 순위

입력:
  sigungu_base.csv    10/14 산출물 (인구·고령화율·인구 증감률·1인당 자가용 승용차)
  sigungu_access.csv  10/15~16 산출물 (최근접 KTX역 접근시간·대도시/서울 승용차 시간)
출력:
  rail_gap.xlsx       스프레드시트 — `가중치`, `시군진단`, `공백상위` 시트 (수식으로 계산, 가중치 바꾸면 순위 갱신)
  sigungu_gap.csv     같은 계산을 파이썬으로 한 결과 (A·지도용)

공백 지수 (노션 1단계 B 지표 5개)
  1) 최근접 KTX역 접근시간      길수록 공백 ↑
  2) 고령화율 2022              높을수록 공백 ↑
  3) 인구 증감률 2013→2022      감소할수록 공백 ↑ (교통약자·지역 쇠퇴)
  4) 최근접 대도시 승용차 시간    길수록 공백 ↑
  5) 1인당 자가용 승용차 2022    많을수록 공백 ↑ (승용차 의존)
  각 지표를 분석 대상 시군 안에서 0~100으로 최소-최대 정규화 → 가중평균 (기본 가중치 같게 0.2)
분석 대상에서 제외(포함=0): 섬(제주시·서귀포시·울릉군 — 카페리 경로), 부산 중구·대구 중구(1인당 등록대수 이상치)
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INDICATORS = [  # (진단표 열 이름, 원자료 열, 방향 +1=클수록 공백, 설명)
    ("KTX역_접근시간_분", "KTX역_접근시간_분", 1, "시군 인구가중 중심 → 최근접 KTX역(A station_coords 69개) 승용차 시간, 평일 10시"),
    ("고령화율_2022", "고령화율_2022", 1, "65세 이상 ÷ 총인구 × 100, 2022.12 주민등록"),
    ("인구증감률_2013_2022", "인구증감률_2013_2022", -1, "(2022 ÷ 2013 − 1) × 100. 감소할수록 공백 점수 ↑"),
    ("대도시_승용차시간_분", "대도시_승용차시간_분", 1, "시군 중심 → 가장 빨리 닿는 특별·광역시청 승용차 시간"),
    ("1인당자가용승용차_2022", "1인당자가용승용차_2022", 1, "자가용 승용차 ÷ 인구, 2022.12"),
]
EXCLUDE = {("제주", "제주시"): "섬(카페리 경로)", ("제주", "서귀포시"): "섬(카페리 경로)",
           ("경북", "울릉군"): "섬(카페리 경로)", ("부산", "중구"): "1인당 등록대수 이상치",
           ("대구", "중구"): "1인당 등록대수 이상치"}
TOP_N = 15

base = pd.read_csv("sigungu_base.csv", dtype={"시군코드": str})
acc = pd.read_csv("sigungu_access.csv", dtype={"시군코드": str})
df = base.merge(acc.drop(columns=["시도", "시군", "비고"]), on="시군코드", how="inner", validate="1:1")
assert len(df) == len(base) == len(acc) == 229, len(df)
df["포함"] = [0 if (s, g) in EXCLUDE else 1 for s, g in zip(df["시도"], df["시군"])]
df["제외사유"] = [EXCLUDE.get((s, g), "") for s, g in zip(df["시도"], df["시군"])]

# ---------- 파이썬 계산 (엑셀 수식과 같은 방식, 검증·CSV용) ----------
inc = df["포함"] == 1
for col, src, sign, _ in INDICATORS:
    lo, hi = df.loc[inc, src].min(), df.loc[inc, src].max()
    s = (df[src] - lo) / (hi - lo) * 100
    df[f"점수_{col}"] = (s if sign > 0 else 100 - s).round(6)
w = [1 / len(INDICATORS)] * len(INDICATORS)
df["공백지수"] = sum(wi * df[f"점수_{c}"] for wi, (c, *_) in zip(w, INDICATORS))
df.loc[~inc, "공백지수"] = None
df["순위"] = df["공백지수"].rank(ascending=False, method="first")

# ---------- 엑셀 ----------
F = "Arial"
font = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
blue = Font(name=F, size=10, color="0000FF")
green = Font(name=F, size=10, color="008000")
head_fill = PatternFill("solid", fgColor="D9E1F2")
input_fill = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

# 가중치 시트
ws = wb.active
ws.title = "가중치"
ws["A1"] = "철도 공백 지수 — 지표·가중치 (노란 칸만 수정)"
ws["A1"].font = Font(name=F, size=12, bold=True)
hdr = ["지표", "방향", "가중치", "설명"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(3, j, h); c.font = bold; c.fill = head_fill; c.border = box
for i, (col, src, sign, desc) in enumerate(INDICATORS, 4):
    ws.cell(i, 1, col).font = font
    d = ws.cell(i, 2, sign); d.font = blue; d.fill = input_fill
    d.comment = Comment("+1 = 값이 클수록 공백 점수 높음, -1 = 값이 작을수록 공백 점수 높음", "B")
    wc = ws.cell(i, 3, 0.2); wc.font = blue; wc.fill = input_fill; wc.number_format = "0.00"
    ws.cell(i, 4, desc).font = font
    for j in range(1, 5):
        ws.cell(i, j).border = box
last = 3 + len(INDICATORS)
ws.cell(last + 1, 1, "합계").font = bold
s = ws.cell(last + 1, 3, f"=SUM(C4:C{last})"); s.font = bold; s.number_format = "0.00"
notes = [
    "정규화: 지표마다 분석 대상(시군진단 시트 포함=1) 안에서 최소-최대 → 0~100점. 방향 -1이면 100 − 점수",
    "공백 지수 = Σ(가중치 × 점수) ÷ Σ가중치. 가중치 합이 1이 아니어도 됨",
    "기본 가중치 0.2(5개 같게)는 B 가정 — 10/19 A 민감도 분석에서 바꿔 볼 값",
    "인구 증감률 방향(-1)은 B 가정: 인구가 줄어드는 지역을 이동 약자·쇠퇴 지역으로 보고 공백 ↑",
    "출처: sigungu_base.csv(KOSIS 주민등록인구·국토부 자동차등록), sigungu_access.csv(카카오모빌리티 길찾기, SGIS 2025.2Q 경계)",
]
for k, n in enumerate(notes):
    ws.cell(last + 3 + k, 1, n).font = font
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 7
ws.column_dimensions["C"].width = 9
ws.column_dimensions["D"].width = 80

# 시군진단 시트
sd = wb.create_sheet("시군진단")
info = ["시군코드", "시도", "시군", "인구_2022", "고령인구_2022"]
ref = ["서울_승용차시간_분", "최근접KTX역", "최근접대도시"]
raw = [c for c, *_ in INDICATORS]
cols = info + raw + ref + ["포함", "제외사유"] + [f"점수_{c}" for c in raw] + ["공백지수", "순위", "순위_고유"]
for j, h in enumerate(cols, 1):
    c = sd.cell(1, j, h); c.font = bold; c.fill = head_fill; c.border = box
    c.alignment = Alignment(wrap_text=True, vertical="center")
L = {h: get_column_letter(j) for j, h in enumerate(cols, 1)}
n = len(df)
R1, RN = 2, n + 1
for i, rd in enumerate(df.to_dict("records"), R1):
    for h in info + raw + ref + ["제외사유"]:
        sd[f"{L[h]}{i}"] = rd[h]
    p = sd[f"{L['포함']}{i}"]; p.value = rd["포함"]; p.font = blue; p.fill = input_fill
    for k, c in enumerate(raw):
        x, wrow = L[c], 4 + k
        rng = f"${x}${R1}:${x}${RN}"
        inc_rng = f"${L['포함']}${R1}:${L['포함']}${RN}"
        lo = f"_xlfn.MINIFS({rng},{inc_rng},1)"
        hi = f"_xlfn.MAXIFS({rng},{inc_rng},1)"
        sc = f"({x}{i}-{lo})/({hi}-{lo})*100"
        sd[f"{L['점수_' + c]}{i}"] = f"=IF(가중치!$B${wrow}>0,{sc},100-{sc})"
    wsum = "+".join(f"가중치!$C${4 + k}*{L['점수_' + c]}{i}" for k, c in enumerate(raw))
    sd[f"{L['공백지수']}{i}"] = f"=IF({L['포함']}{i}=1,({wsum})/가중치!$C${last + 1},\"\")"
    gi = f"${L['공백지수']}${R1}:${L['공백지수']}${RN}"
    sd[f"{L['순위']}{i}"] = f"=IF({L['포함']}{i}=1,COUNTIF({gi},\">\"&{L['공백지수']}{i})+1,\"제외\")"
    sd[f"{L['순위_고유']}{i}"] = (f"=IF({L['포함']}{i}=1,{L['순위']}{i}+COUNTIF(${L['공백지수']}${R1}:{L['공백지수']}{i},"
                                f"{L['공백지수']}{i})-1,\"\")")
for i in range(R1, RN + 1):
    for j in range(1, len(cols) + 1):
        c = sd.cell(i, j)
        if c.font != blue:
            c.font = font
        c.border = box
    for h in raw + ["서울_승용차시간_분"]:
        sd[f"{L[h]}{i}"].number_format = "0.0000" if h.startswith("1인당") else "0.0"
    for h in [f"점수_{c}" for c in raw] + ["공백지수"]:
        sd[f"{L[h]}{i}"].number_format = "0.0"
    for h in ["인구_2022", "고령인구_2022"]:
        sd[f"{L[h]}{i}"].number_format = "#,##0"
sd.freeze_panes = "D2"
sd.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{RN}"
for j, h in enumerate(cols, 1):
    sd.column_dimensions[get_column_letter(j)].width = 9 if h.startswith("점수_") else max(8, min(16, len(h) + 2))
sd.row_dimensions[1].height = 45
sd[f"{L['포함']}1"].comment = Comment("1 = 공백 지수 계산 대상, 0 = 제외(섬·이상치). 바꾸면 정규화 범위·순위가 다시 계산됨", "B")
sd[f"{L['순위_고유']}1"].comment = Comment("동점일 때 위쪽 행을 앞 순위로 둔 고유 순위 (공백상위 시트 조회용)", "B")

# 공백상위 시트
tp = wb.create_sheet("공백상위")
tp["A1"] = f"철도 공백 지수 상위 {TOP_N}개 시군 (시군진단 시트 수식에서 자동 조회)"
tp["A1"].font = Font(name=F, size=12, bold=True)
show = ["시도", "시군", "공백지수"] + raw + ["최근접KTX역", "최근접대도시", "인구_2022"]
tp.cell(3, 1, "순위").font = bold
tp.cell(3, 1).fill = head_fill
for j, h in enumerate(show, 2):
    c = tp.cell(3, j, h); c.font = bold; c.fill = head_fill; c.border = box
    c.alignment = Alignment(wrap_text=True, vertical="center")
uniq = f"시군진단!${L['순위_고유']}${R1}:${L['순위_고유']}${RN}"
for k in range(1, TOP_N + 1):
    row = 3 + k
    tp.cell(row, 1, k).font = font
    for j, h in enumerate(show, 2):
        c = tp.cell(row, j, f"=INDEX(시군진단!${L[h]}${R1}:${L[h]}${RN},MATCH($A{row},{uniq},0))")
        c.font = green
        c.border = box
        c.number_format = "0.0000" if h.startswith("1인당") else ("#,##0" if h == "인구_2022" else "0.0")
tp.row_dimensions[3].height = 45
for j in range(1, len(show) + 2):
    tp.column_dimensions[get_column_letter(j)].width = 12
tp.column_dimensions["A"].width = 6

wb.save("rail_gap.xlsx")

out = df[info + raw + ref + ["포함", "제외사유"] + [f"점수_{c}" for c in raw] + ["공백지수", "순위"]].copy()
out = out.round({**{f"점수_{c}": 1 for c in raw}, "공백지수": 2})
out["순위"] = out["순위"].astype("Int64")
out.sort_values(["순위", "시군코드"], na_position="last").to_csv("sigungu_gap.csv", index=False, encoding="utf-8-sig")
print(f"시군 {len(df)}개, 분석 대상 {inc.sum()}개, 제외 {(~inc).sum()}개")
print(out.nsmallest(TOP_N, "순위")[["순위", "시도", "시군", "공백지수"] + raw].to_string(index=False))
